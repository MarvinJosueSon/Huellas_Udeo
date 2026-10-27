import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime, date, timedelta
from decimal import Decimal

try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        Font,
        PatternFill,
        Alignment,
        Border,
        Side
    )
    from openpyxl.utils import get_column_letter

    OPENPYXL_DISPONIBLE = True

except ImportError:

    Workbook = None
    Font = None
    PatternFill = None
    Alignment = None
    Border = None
    Side = None
    get_column_letter = None

    OPENPYXL_DISPONIBLE = False


class Reportes(tk.Frame):

    def __init__(
        self,
        master,
        database
    ):
        super().__init__(
            master,
            bg="#F3F3F3"
        )

        self.database = database

        self.mapa_secciones = {}
        self.mapa_periodos = {}

        self.columnas_reporte_actual = []
        self.datos_reporte_actual = []
        self.titulo_reporte_actual = ""
        self.detalles_reporte_actual = []

        self.pack(
            fill="both",
            expand=True
        )

        self.crear_interfaz()
        self.cargar_secciones()
        self.cargar_periodos()
        self.cambiar_tipo_reporte()


    def crear_interfaz(self):

        encabezado = tk.Frame(
            self,
            bg="#F3F3F3"
        )

        encabezado.pack(
            fill="x",
            padx=30,
            pady=(25, 15)
        )


        tk.Label(
            encabezado,
            text="Reportes",
            bg="#F3F3F3",
            fg="#202020",
            font=("Segoe UI", 24, "bold")
        ).pack(
            side="left"
        )


        self.btn_exportar = tk.Button(
            encabezado,
            text="Exportar a Excel",
            bg="#107C10",
            fg="#FFFFFF",
            activebackground="#0B6A0B",
            activeforeground="#FFFFFF",
            relief="flat",
            bd=0,
            cursor="hand2",
            font=("Segoe UI", 10, "bold"),
            command=self.exportar_excel,
            state="disabled"
        )

        self.btn_exportar.pack(
            side="right",
            ipadx=15,
            ipady=7
        )


        panel_control = tk.Frame(
            self,
            bg="#FFFFFF",
            highlightthickness=1,
            highlightbackground="#DDDDDD"
        )

        panel_control.pack(
            fill="x",
            padx=30,
            pady=(0, 15)
        )


        tk.Label(
            panel_control,
            text="Tipo de reporte",
            bg="#FFFFFF",
            fg="#202020",
            font=("Segoe UI", 10, "bold")
        ).grid(
            row=0,
            column=0,
            sticky="w",
            padx=(20, 10),
            pady=(15, 5)
        )


        self.cmb_tipo = ttk.Combobox(
            panel_control,
            state="readonly",
            width=35,
            values=(
                "Asistencia por curso/sección",
                "Puntualidad",
                "Asistencia por fecha",
                "Justificaciones",
                "Final de periodo"
            )
        )

        self.cmb_tipo.grid(
            row=1,
            column=0,
            sticky="ew",
            padx=(20, 10),
            pady=(0, 15)
        )

        self.cmb_tipo.current(
            0
        )

        self.cmb_tipo.bind(
            "<<ComboboxSelected>>",
            self.cambiar_tipo_reporte
        )


        self.panel_parametro = tk.Frame(
            panel_control,
            bg="#FFFFFF"
        )

        self.panel_parametro.grid(
            row=0,
            column=1,
            rowspan=2,
            sticky="nsew",
            padx=10,
            pady=10
        )


        tk.Button(
            panel_control,
            text="Consultar",
            bg="#0078D4",
            fg="#FFFFFF",
            activebackground="#106EBE",
            activeforeground="#FFFFFF",
            relief="flat",
            bd=0,
            cursor="hand2",
            font=("Segoe UI", 10, "bold"),
            command=self.consultar_reporte
        ).grid(
            row=1,
            column=2,
            padx=(10, 20),
            pady=(0, 15),
            ipadx=20,
            ipady=5
        )


        panel_control.grid_columnconfigure(
            0,
            weight=1
        )

        panel_control.grid_columnconfigure(
            1,
            weight=2
        )


        panel_info = tk.Frame(
            self,
            bg="#F3F3F3"
        )

        panel_info.pack(
            fill="x",
            padx=30,
            pady=(0, 8)
        )


        self.lbl_resultados = tk.Label(
            panel_info,
            text="Resultados: 0",
            bg="#F3F3F3",
            fg="#666666",
            font=("Segoe UI", 9)
        )

        self.lbl_resultados.pack(
            side="left"
        )


        panel_tabla = tk.Frame(
            self,
            bg="#FFFFFF",
            highlightthickness=1,
            highlightbackground="#DDDDDD"
        )

        panel_tabla.pack(
            fill="both",
            expand=True,
            padx=30,
            pady=(0, 20)
        )


        self.tabla = ttk.Treeview(
            panel_tabla,
            show="headings",
            selectmode="browse"
        )


        scroll_vertical = ttk.Scrollbar(
            panel_tabla,
            orient="vertical",
            command=self.tabla.yview
        )


        scroll_horizontal = ttk.Scrollbar(
            panel_tabla,
            orient="horizontal",
            command=self.tabla.xview
        )


        self.tabla.configure(
            yscrollcommand=scroll_vertical.set,
            xscrollcommand=scroll_horizontal.set
        )


        self.tabla.grid(
            row=0,
            column=0,
            sticky="nsew"
        )

        scroll_vertical.grid(
            row=0,
            column=1,
            sticky="ns"
        )

        scroll_horizontal.grid(
            row=1,
            column=0,
            sticky="ew"
        )


        panel_tabla.grid_rowconfigure(
            0,
            weight=1
        )

        panel_tabla.grid_columnconfigure(
            0,
            weight=1
        )


    def cargar_secciones(self):

        resultados = (
            self.database
            .llamar_procedimiento(
                "ListarSeccionesDocente"
            )
        )


        if resultados is None:

            messagebox.showerror(
                "Reportes",
                (
                    "No se pudieron cargar "
                    "las secciones.\n\n"
                    + str(
                        self.database
                        .obtener_ultimo_error()
                    )
                )
            )

            return


        secciones = []

        if resultados:

            secciones = resultados[0]


        valores = []

        self.mapa_secciones = {}


        for seccion in secciones:

            texto = (
                f"{seccion['id']} - "
                f"{seccion['curso']} - "
                f"Sección {seccion['seccion']} - "
                f"{seccion['periodo']} - "
                f"{seccion['jornada']}"
            )

            valores.append(
                texto
            )

            self.mapa_secciones[
                texto
            ] = seccion["id"]


        self.valores_secciones = valores


    def cargar_periodos(self):

        resultados = (
            self.database
            .llamar_procedimiento(
                "ListarPeriodos"
            )
        )


        if resultados is None:

            messagebox.showerror(
                "Reportes",
                (
                    "No se pudieron cargar "
                    "los periodos.\n\n"
                    + str(
                        self.database
                        .obtener_ultimo_error()
                    )
                )
            )

            return


        periodos = []

        if resultados:

            periodos = resultados[0]


        valores = []

        self.mapa_periodos = {}


        for periodo in periodos:

            texto = (
                f"{periodo['id']} - "
                f"{periodo['nombre']} "
                f"({periodo['estado']})"
            )

            valores.append(
                texto
            )

            self.mapa_periodos[
                texto
            ] = periodo["id"]


        self.valores_periodos = valores


    def limpiar_parametros(self):

        for widget in (
            self.panel_parametro
            .winfo_children()
        ):

            widget.destroy()


    def cambiar_tipo_reporte(
        self,
        event=None
    ):

        self.limpiar_parametros()
        self.limpiar_tabla()

        tipo = (
            self.cmb_tipo
            .get()
        )


        if tipo in (
            "Asistencia por curso/sección",
            "Puntualidad"
        ):

            tk.Label(
                self.panel_parametro,
                text="Sección",
                bg="#FFFFFF",
                fg="#202020",
                font=("Segoe UI", 9)
            ).pack(
                anchor="w",
                pady=(0, 5)
            )


            self.cmb_seccion = ttk.Combobox(
                self.panel_parametro,
                state="readonly",
                width=55,
                values=self.valores_secciones
            )

            self.cmb_seccion.pack(
                fill="x"
            )


            if self.valores_secciones:

                self.cmb_seccion.current(
                    0
                )


        elif tipo == "Asistencia por fecha":

            tk.Label(
                self.panel_parametro,
                text="Fecha (AAAA-MM-DD)",
                bg="#FFFFFF",
                fg="#202020",
                font=("Segoe UI", 9)
            ).pack(
                anchor="w",
                pady=(0, 5)
            )


            self.txt_fecha = ttk.Entry(
                self.panel_parametro,
                width=25
            )

            self.txt_fecha.pack(
                fill="x"
            )


            self.txt_fecha.insert(
                0,
                datetime.now().strftime(
                    "%Y-%m-%d"
                )
            )


        elif tipo == "Justificaciones":

            tk.Label(
                self.panel_parametro,
                text=(
                    "Se mostrarán todas las "
                    "justificaciones permitidas "
                    "para tu usuario."
                ),
                bg="#FFFFFF",
                fg="#666666",
                font=("Segoe UI", 9)
            ).pack(
                anchor="w",
                pady=(15, 0)
            )


        elif tipo == "Final de periodo":

            tk.Label(
                self.panel_parametro,
                text="Periodo académico",
                bg="#FFFFFF",
                fg="#202020",
                font=("Segoe UI", 9)
            ).pack(
                anchor="w",
                pady=(0, 5)
            )


            self.cmb_periodo = ttk.Combobox(
                self.panel_parametro,
                state="readonly",
                width=45,
                values=self.valores_periodos
            )

            self.cmb_periodo.pack(
                fill="x"
            )


            if self.valores_periodos:

                self.cmb_periodo.current(
                    0
                )


    def consultar_reporte(self):

        tipo = (
            self.cmb_tipo
            .get()
        )


        if tipo == "Asistencia por curso/sección":

            self.reporte_asistencia_curso()


        elif tipo == "Puntualidad":

            self.reporte_puntualidad()


        elif tipo == "Asistencia por fecha":

            self.reporte_asistencia_fecha()


        elif tipo == "Justificaciones":

            self.reporte_justificaciones()


        elif tipo == "Final de periodo":

            self.reporte_final_periodo()


    def reporte_asistencia_curso(self):

        seccion_texto = (
            self.cmb_seccion
            .get()
            .strip()
        )


        if (
            seccion_texto
            not in self.mapa_secciones
        ):

            messagebox.showwarning(
                "Reportes",
                "Selecciona una sección."
            )

            return


        id_seccion = (
            self.mapa_secciones[
                seccion_texto
            ]
        )


        resultados = (
            self.database
            .llamar_procedimiento(
                "ReporteAsistenciaCurso",
                (
                    id_seccion,
                )
            )
        )


        if not self.validar_resultado(
            resultados
        ):

            return


        datos = self.obtener_datos(
            resultados
        )


        columnas = [
            (
                "codigo",
                "Código",
                120
            ),
            (
                "alumno",
                "Alumno",
                220
            ),
            (
                "curso",
                "Curso",
                180
            ),
            (
                "seccion",
                "Sección",
                80
            ),
            (
                "total_clases",
                "Clases",
                80
            ),
            (
                "presentes",
                "Presentes",
                90
            ),
            (
                "tardes",
                "Tardes",
                80
            ),
            (
                "justificadas",
                "Justificadas",
                100
            ),
            (
                "ausentes",
                "Ausentes",
                90
            ),
            (
                "fuera_tiempo",
                "Fuera tiempo",
                110
            ),
            (
                "porcentaje",
                "Porcentaje",
                100
            )
        ]


        self.mostrar_resultado(
            columnas,
            datos
        )


    def reporte_puntualidad(self):

        seccion_texto = (
            self.cmb_seccion
            .get()
            .strip()
        )


        if (
            seccion_texto
            not in self.mapa_secciones
        ):

            messagebox.showwarning(
                "Reportes",
                "Selecciona una sección."
            )

            return


        id_seccion = (
            self.mapa_secciones[
                seccion_texto
            ]
        )


        resultados = (
            self.database
            .llamar_procedimiento(
                "ReportePuntualidad",
                (
                    id_seccion,
                )
            )
        )


        if not self.validar_resultado(
            resultados
        ):

            return


        datos = self.obtener_datos(
            resultados
        )


        columnas = [
            (
                "codigo",
                "Código",
                120
            ),
            (
                "alumno",
                "Alumno",
                240
            ),
            (
                "curso",
                "Curso",
                200
            ),
            (
                "seccion",
                "Sección",
                90
            ),
            (
                "puntuales",
                "Puntuales",
                100
            ),
            (
                "tardanzas",
                "Tardanzas",
                100
            ),
            (
                "porcentaje_puntualidad",
                "% Puntualidad",
                120
            )
        ]


        self.mostrar_resultado(
            columnas,
            datos
        )


    def reporte_asistencia_fecha(self):

        fecha_texto = (
            self.txt_fecha
            .get()
            .strip()
        )


        try:

            datetime.strptime(
                fecha_texto,
                "%Y-%m-%d"
            )

        except ValueError:

            messagebox.showwarning(
                "Reportes",
                (
                    "La fecha debe utilizar "
                    "el formato AAAA-MM-DD."
                )
            )

            return


        resultados = (
            self.database
            .llamar_procedimiento(
                "ReporteAsistenciaFecha",
                (
                    fecha_texto,
                )
            )
        )


        if not self.validar_resultado(
            resultados
        ):

            return


        datos = self.obtener_datos(
            resultados
        )


        columnas = [
            (
                "codigo",
                "Código",
                110
            ),
            (
                "alumno",
                "Alumno",
                210
            ),
            (
                "curso",
                "Curso",
                180
            ),
            (
                "seccion",
                "Sección",
                80
            ),
            (
                "jornada",
                "Jornada",
                110
            ),
            (
                "fecha",
                "Fecha",
                100
            ),
            (
                "hora_clase",
                "Hora clase",
                100
            ),
            (
                "hora_registro",
                "Registro",
                160
            ),
            (
                "estado",
                "Estado",
                120
            )
        ]


        self.mostrar_resultado(
            columnas,
            datos
        )


    def reporte_justificaciones(self):

        resultados = (
            self.database
            .llamar_procedimiento(
                "ReporteJustificaciones"
            )
        )


        if not self.validar_resultado(
            resultados
        ):

            return


        datos = self.obtener_datos(
            resultados
        )


        columnas = [
            (
                "id",
                "ID",
                55
            ),
            (
                "codigo",
                "Código",
                110
            ),
            (
                "alumno",
                "Alumno",
                200
            ),
            (
                "curso",
                "Curso",
                170
            ),
            (
                "seccion",
                "Sección",
                80
            ),
            (
                "motivo",
                "Motivo",
                260
            ),
            (
                "estado",
                "Estado",
                100
            ),
            (
                "fecha_solicitud",
                "Solicitud",
                160
            ),
            (
                "fecha_respuesta",
                "Respuesta",
                160
            )
        ]


        self.mostrar_resultado(
            columnas,
            datos
        )


    def reporte_final_periodo(self):

        periodo_texto = (
            self.cmb_periodo
            .get()
            .strip()
        )


        if (
            periodo_texto
            not in self.mapa_periodos
        ):

            messagebox.showwarning(
                "Reportes",
                (
                    "Selecciona un "
                    "periodo académico."
                )
            )

            return


        id_periodo = (
            self.mapa_periodos[
                periodo_texto
            ]
        )


        resultados = (
            self.database
            .llamar_procedimiento(
                "ReporteFinalPeriodo",
                (
                    id_periodo,
                )
            )
        )


        if not self.validar_resultado(
            resultados
        ):

            return


        datos = self.obtener_datos(
            resultados
        )


        columnas = [
            (
                "codigo",
                "Código",
                110
            ),
            (
                "alumno",
                "Alumno",
                200
            ),
            (
                "curso",
                "Curso",
                170
            ),
            (
                "seccion",
                "Sección",
                80
            ),
            (
                "periodo",
                "Periodo",
                160
            ),
            (
                "total_clases",
                "Clases",
                80
            ),
            (
                "presentes",
                "Presentes",
                90
            ),
            (
                "tardes",
                "Tardes",
                80
            ),
            (
                "justificadas",
                "Justificadas",
                100
            ),
            (
                "ausentes",
                "Ausentes",
                90
            ),
            (
                "fuera_tiempo",
                "Fuera tiempo",
                110
            ),
            (
                "porcentaje",
                "Porcentaje",
                100
            )
        ]


        self.mostrar_resultado(
            columnas,
            datos
        )


    def validar_resultado(
        self,
        resultados
    ):

        if resultados is None:

            messagebox.showerror(
                "Reportes",
                (
                    "No se pudo generar "
                    "el reporte.\n\n"
                    + str(
                        self.database
                        .obtener_ultimo_error()
                    )
                )
            )

            return False


        return True


    def obtener_datos(
        self,
        resultados
    ):

        if not resultados:

            return []


        return resultados[0]


    def mostrar_resultado(
        self,
        columnas,
        datos
    ):

        self.configurar_tabla(
            columnas
        )


        self.columnas_reporte_actual = columnas
        self.datos_reporte_actual = datos
        self.titulo_reporte_actual = (
            self.cmb_tipo.get()
        )

        self.detalles_reporte_actual = (
            self.obtener_detalles_exportacion()
        )


        if datos:

            self.btn_exportar.config(
                state="normal"
            )

        else:

            self.btn_exportar.config(
                state="disabled"
            )


        for registro in datos:

            valores = []


            for clave, titulo, ancho in columnas:

                valor = registro.get(
                    clave,
                    ""
                )

                valores.append(
                    self.formatear_valor(
                        clave,
                        valor
                    )
                )


            self.tabla.insert(
                "",
                tk.END,
                values=valores
            )


        self.lbl_resultados.config(
            text=(
                f"Resultados: "
                f"{len(datos)}"
            )
        )


    def configurar_tabla(
        self,
        columnas
    ):

        self.limpiar_tabla()


        claves = [
            columna[0]
            for columna in columnas
        ]


        self.tabla[
            "columns"
        ] = claves


        for clave, titulo, ancho in columnas:

            self.tabla.heading(
                clave,
                text=titulo
            )

            self.tabla.column(
                clave,
                width=ancho,
                minwidth=60,
                anchor=(
                    "center"
                    if clave in (
                        "seccion",
                        "total_clases",
                        "presentes",
                        "tardes",
                        "justificadas",
                        "ausentes",
                        "fuera_tiempo",
                        "puntuales",
                        "tardanzas",
                        "porcentaje",
                        "porcentaje_puntualidad",
                        "estado"
                    )
                    else "w"
                )
            )


    def limpiar_tabla(self):

        for fila in (
            self.tabla.get_children()
        ):

            self.tabla.delete(
                fila
            )


        self.lbl_resultados.config(
            text="Resultados: 0"
        )


        self.columnas_reporte_actual = []
        self.datos_reporte_actual = []
        self.titulo_reporte_actual = ""
        self.detalles_reporte_actual = []


        if hasattr(
            self,
            "btn_exportar"
        ):

            self.btn_exportar.config(
                state="disabled"
            )


    def obtener_detalles_exportacion(self):

        tipo = (
            self.cmb_tipo
            .get()
        )

        detalles = []


        if tipo in (
            "Asistencia por curso/sección",
            "Puntualidad"
        ):

            if hasattr(
                self,
                "cmb_seccion"
            ):

                seccion = (
                    self.cmb_seccion
                    .get()
                    .strip()
                )

                if seccion:

                    detalles.append(
                        (
                            "Sección",
                            seccion
                        )
                    )


        elif tipo == "Asistencia por fecha":

            if hasattr(
                self,
                "txt_fecha"
            ):

                fecha = (
                    self.txt_fecha
                    .get()
                    .strip()
                )

                if fecha:

                    detalles.append(
                        (
                            "Fecha",
                            fecha
                        )
                    )


        elif tipo == "Final de periodo":

            if hasattr(
                self,
                "cmb_periodo"
            ):

                periodo = (
                    self.cmb_periodo
                    .get()
                    .strip()
                )

                if periodo:

                    detalles.append(
                        (
                            "Periodo",
                            periodo
                        )
                    )


        return detalles


    def exportar_excel(self):

        if not OPENPYXL_DISPONIBLE:

            messagebox.showerror(
                "Exportar a Excel",
                (
                    "No se encontró la librería "
                    "openpyxl.\n\n"
                    "Instálala con:\n"
                    "pip install openpyxl"
                )
            )

            return


        if not self.datos_reporte_actual:

            messagebox.showwarning(
                "Exportar a Excel",
                (
                    "Primero debes consultar "
                    "un reporte con resultados."
                )
            )

            return


        nombre_reporte = (
            self.titulo_reporte_actual
            .lower()
            .replace(
                "/",
                "_"
            )
            .replace(
                " ",
                "_"
            )
            .replace(
                "ó",
                "o"
            )
            .replace(
                "í",
                "i"
            )
            .replace(
                "á",
                "a"
            )
            .replace(
                "é",
                "e"
            )
            .replace(
                "ú",
                "u"
            )
        )


        nombre_archivo = (
            f"{nombre_reporte}_"
            f"{datetime.now().strftime('%Y%m%d')}"
            ".xlsx"
        )


        ruta = filedialog.asksaveasfilename(
            title="Guardar reporte en Excel",
            defaultextension=".xlsx",
            initialfile=nombre_archivo,
            filetypes=[
                (
                    "Archivo de Excel",
                    "*.xlsx"
                )
            ]
        )


        if not ruta:

            return


        try:

            libro = Workbook()

            hoja = libro.active

            hoja.title = "Reporte"


            cantidad_columnas = len(
                self.columnas_reporte_actual
            )


            ultima_columna = (
                get_column_letter(
                    cantidad_columnas
                )
            )


            hoja.merge_cells(
                f"A1:{ultima_columna}1"
            )


            celda_titulo = hoja["A1"]

            celda_titulo.value = (
                "UDEO Assistance Digital"
            )

            celda_titulo.font = Font(
                name="Segoe UI",
                size=16,
                bold=True
            )

            celda_titulo.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )


            hoja.merge_cells(
                f"A2:{ultima_columna}2"
            )


            celda_reporte = hoja["A2"]

            celda_reporte.value = (
                self.titulo_reporte_actual
            )

            celda_reporte.font = Font(
                name="Segoe UI",
                size=13,
                bold=True
            )

            celda_reporte.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )


            fila_actual = 4


            for nombre, valor in (
                self.detalles_reporte_actual
            ):

                hoja.cell(
                    row=fila_actual,
                    column=1,
                    value=nombre
                ).font = Font(
                    bold=True
                )


                hoja.cell(
                    row=fila_actual,
                    column=2,
                    value=valor
                )


                fila_actual += 1


            hoja.cell(
                row=fila_actual,
                column=1,
                value="Fecha de generación"
            ).font = Font(
                bold=True
            )


            hoja.cell(
                row=fila_actual,
                column=2,
                value=datetime.now()
            ).number_format = (
                "yyyy-mm-dd hh:mm:ss"
            )


            fila_actual += 2

            fila_encabezado = fila_actual


            relleno_encabezado = PatternFill(
                fill_type="solid",
                fgColor="0078D4"
            )


            fuente_encabezado = Font(
                color="FFFFFF",
                bold=True
            )


            borde = Border(
                left=Side(
                    style="thin",
                    color="D9D9D9"
                ),
                right=Side(
                    style="thin",
                    color="D9D9D9"
                ),
                top=Side(
                    style="thin",
                    color="D9D9D9"
                ),
                bottom=Side(
                    style="thin",
                    color="D9D9D9"
                )
            )


            for indice, (
                clave,
                titulo,
                ancho
            ) in enumerate(
                self.columnas_reporte_actual,
                start=1
            ):

                celda = hoja.cell(
                    row=fila_encabezado,
                    column=indice,
                    value=titulo
                )


                celda.fill = (
                    relleno_encabezado
                )

                celda.font = (
                    fuente_encabezado
                )

                celda.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )

                celda.border = borde


            fila_datos = (
                fila_encabezado + 1
            )


            for registro in (
                self.datos_reporte_actual
            ):

                for indice, (
                    clave,
                    titulo,
                    ancho
                ) in enumerate(
                    self.columnas_reporte_actual,
                    start=1
                ):

                    valor = registro.get(
                        clave,
                        ""
                    )


                    valor_excel = (
                        self.preparar_valor_excel(
                            valor
                        )
                    )


                    celda = hoja.cell(
                        row=fila_datos,
                        column=indice,
                        value=valor_excel
                    )


                    celda.border = borde

                    celda.alignment = Alignment(
                        vertical="center"
                    )


                    if isinstance(
                        valor,
                        datetime
                    ):

                        celda.number_format = (
                            "yyyy-mm-dd hh:mm:ss"
                        )


                    elif isinstance(
                        valor,
                        date
                    ):

                        celda.number_format = (
                            "yyyy-mm-dd"
                        )


                    if clave in (
                        "porcentaje",
                        "porcentaje_puntualidad"
                    ):

                        celda.number_format = (
                            '0.00"%"'
                        )


                fila_datos += 1


            ultima_fila = (
                fila_datos - 1
            )


            hoja.auto_filter.ref = (
                f"A{fila_encabezado}:"
                f"{ultima_columna}"
                f"{ultima_fila}"
            )


            hoja.freeze_panes = (
                f"A{fila_encabezado + 1}"
            )


            self.ajustar_columnas_excel(
                hoja,
                fila_encabezado
            )


            hoja.row_dimensions[
                1
            ].height = 25

            hoja.row_dimensions[
                2
            ].height = 22


            libro.save(
                ruta
            )


            messagebox.showinfo(
                "Exportar a Excel",
                (
                    "El reporte fue exportado "
                    "correctamente."
                )
            )


        except PermissionError:

            messagebox.showerror(
                "Exportar a Excel",
                (
                    "No se pudo guardar el archivo.\n\n"
                    "Verifica que el archivo no esté "
                    "abierto en Excel."
                )
            )


        except Exception as error:

            messagebox.showerror(
                "Exportar a Excel",
                (
                    "No se pudo exportar "
                    "el reporte.\n\n"
                    f"{error}"
                )
            )


    def preparar_valor_excel(
        self,
        valor
    ):

        if valor is None:

            return ""


        if isinstance(
            valor,
            Decimal
        ):

            return float(
                valor
            )


        if isinstance(
            valor,
            timedelta
        ):

            segundos = int(
                valor.total_seconds()
            )

            horas = (
                segundos // 3600
            )

            minutos = (
                segundos % 3600
            ) // 60

            segundos_resto = (
                segundos % 60
            )


            return (
                f"{horas:02d}:"
                f"{minutos:02d}:"
                f"{segundos_resto:02d}"
            )


        return valor


    def ajustar_columnas_excel(
        self,
        hoja,
        fila_encabezado
    ):

        for indice, (
            clave,
            titulo,
            ancho_tabla
        ) in enumerate(
            self.columnas_reporte_actual,
            start=1
        ):

            letra = (
                get_column_letter(
                    indice
                )
            )


            ancho_maximo = len(
                str(
                    titulo
                )
            )


            for fila in range(
                fila_encabezado,
                hoja.max_row + 1
            ):

                valor = hoja.cell(
                    row=fila,
                    column=indice
                ).value


                if valor is None:

                    continue


                ancho_maximo = max(
                    ancho_maximo,
                    len(
                        str(
                            valor
                        )
                    )
                )


            hoja.column_dimensions[
                letra
            ].width = min(
                max(
                    ancho_maximo + 2,
                    12
                ),
                35
            )


    def formatear_valor(
        self,
        clave,
        valor
    ):

        if valor is None:

            return ""


        if isinstance(
            valor,
            datetime
        ):

            return valor.strftime(
                "%Y-%m-%d %H:%M:%S"
            )


        if isinstance(
            valor,
            date
        ):

            return valor.strftime(
                "%Y-%m-%d"
            )


        if isinstance(
            valor,
            timedelta
        ):

            segundos = int(
                valor.total_seconds()
            )

            horas = (
                segundos // 3600
            )

            minutos = (
                segundos % 3600
            ) // 60

            segundos_resto = (
                segundos % 60
            )


            return (
                f"{horas:02d}:"
                f"{minutos:02d}:"
                f"{segundos_resto:02d}"
            )


        if clave in (
            "porcentaje",
            "porcentaje_puntualidad"
        ):

            return (
                f"{valor}%"
            )


        return str(
            valor
        )